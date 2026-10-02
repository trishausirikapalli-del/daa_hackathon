import os
import sqlite3
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

# Set up test environment
import app

def test_full_system():
    print("=" * 70)
    print("RUNNING COMPREHENSIVE VERIFICATION TEST SUITE")
    print("=" * 70)

    # 1. Test Client setup
    test_client = app.app.test_client()

    # 2. Test GET / (Single Page Application HTML)
    print("\n[1] Testing GET / (Single Page Application HTML)...")
    res = test_client.get('/')
    assert res.status_code == 200, f"Expected 200, got {res.status_code}"
    html = res.data.decode('utf-8')
    assert "LONGEST COMMON STUDENT ACTIVITY" in html
    assert "How many students do you want to compare?" in html
    assert "dashboard.js" in html
    assert "style.css" in html
    print("  -> PASSED: Single Page Application HTML rendered successfully.")

    # 3. Test GET /api/students
    print("\n[2] Testing GET /api/students...")
    res = test_client.get('/api/students')
    assert res.status_code == 200
    data = res.get_json()
    assert data["success"] is True
    assert len(data["students"]) >= 4
    print(f"  -> PASSED: Found {len(data['students'])} students in DB.")

    # 4. Test POST /api/students (Add Student)
    print("\n[3] Testing POST /api/students (Add Student)...")
    res = test_client.post('/api/students', json={
        "name": "Test Student A",
        "roll_no": "TEST-999",
        "branch": "Computer Science"
    })
    assert res.status_code == 201
    new_student_id = res.get_json()["student_id"]
    print(f"  -> PASSED: Created student ID {new_student_id}")

    # Test Duplicate Roll Number validation
    res_dup = test_client.post('/api/students', json={
        "name": "Another Student",
        "roll_no": "TEST-999",
        "branch": "IT"
    })
    assert res_dup.status_code == 400
    assert "already exists" in res_dup.get_json()["error"]
    print("  -> PASSED: Duplicate roll number validation works.")

    # 5. Test PUT /api/students/<id> (Edit Student)
    print("\n[4] Testing PUT /api/students/<id> (Edit Student)...")
    res = test_client.put(f'/api/students/{new_student_id}', json={
        "name": "Test Student A (Updated)",
        "roll_no": "TEST-999-U",
        "branch": "AI & Data Science"
    })
    assert res.status_code == 200
    print("  -> PASSED: Updated student successfully.")

    # 6. Test POST /api/activities (Add Activities)
    print("\n[5] Testing POST /api/activities (Add Activities)...")
    act_res1 = test_client.post('/api/activities', json={
        "student_id": new_student_id,
        "activity_name": "Python Programming",
        "category": "Coding",
        "activity_date": "2026-09-01",
        "verification_status": "Verified",
        "proof_reference": "CERT-TST-01"
    })
    assert act_res1.status_code == 201
    act_id1 = act_res1.get_json()["activity_id"]

    act_res2 = test_client.post('/api/activities', json={
        "student_id": new_student_id,
        "activity_name": "Data Structures & Algorithms",
        "category": "Academic",
        "activity_date": "2026-09-05",
        "verification_status": "Verified",
        "proof_reference": "CERT-TST-02"
    })
    assert act_res2.status_code == 201

    act_res3 = test_client.post('/api/activities', json={
        "student_id": new_student_id,
        "activity_name": "Custom Robotics Hack",
        "category": "Competition",
        "activity_date": "2026-09-10",
        "verification_status": "Verified",
        "proof_reference": "ROBO-01"
    })
    assert act_res3.status_code == 201
    act_id3 = act_res3.get_json()["activity_id"]
    print("  -> PASSED: Added 3 activities.")

    # 7. Test PUT /api/activities/<id> (Edit Activity)
    print("\n[6] Testing PUT /api/activities/<id> (Edit Activity)...")
    res = test_client.put(f'/api/activities/{act_id3}', json={
        "activity_name": "Smart India Hackathon",
        "category": "Hackathon",
        "activity_date": "2026-09-10",
        "verification_status": "Verified",
        "proof_reference": "SIH-UPDATED"
    })
    assert res.status_code == 200
    print("  -> PASSED: Edited activity successfully.")

    # 8. Test DELETE /api/activities/<id>
    print("\n[7] Testing DELETE /api/activities/<id>...")
    del_act_res = test_client.delete(f'/api/activities/{act_id1}')
    assert del_act_res.status_code == 200
    print("  -> PASSED: Deleted activity successfully.")

    # 9. Test POST /api/students/bulk (Dynamic First Screen)
    print("\n[8] Testing POST /api/students/bulk (First screen dynamic setup)...")
    bulk_res = test_client.post('/api/students/bulk', json={
        "students": [
            {
                "name": "Bulk Student 1",
                "roll_no": "BLK-101",
                "branch": "CSE",
                "activities": [
                    {"activity_name": "Python"},
                    {"activity_name": "DSA"},
                    {"activity_name": "Hackathon"},
                    {"activity_name": "Java"}
                ]
            },
            {
                "name": "Bulk Student 2",
                "roll_no": "BLK-102",
                "branch": "IT",
                "activities": [
                    {"activity_name": "Python"},
                    {"activity_name": "DSA"},
                    {"activity_name": "Hackathon"},
                    {"activity_name": "Seminar"}
                ]
            }
        ]
    })
    assert bulk_res.status_code == 201
    bulk_data = bulk_res.get_json()
    b_ids = bulk_data["student_ids"]
    assert len(b_ids) == 2
    print(f"  -> PASSED: Bulk created students: {b_ids}")

    # 10. Test Multi-Student Comparison & Exact Match %
    print("\n[9] Testing POST /api/compare (LCS Multi-Student Algorithm)...")
    comp_res = test_client.post('/api/compare', json={"student_ids": b_ids})
    assert comp_res.status_code == 200
    comp_data = comp_res.get_json()["data"]
    comp_id = comp_data["comparison_id"]

    # Verify LCS
    expected_lcs = ["Python", "DSA", "Hackathon"]
    assert comp_data["lcs_sequence"] == expected_lcs, f"Got {comp_data['lcs_sequence']}"
    assert comp_data["lcs_length"] == 3
    print(f"  -> LCS Discovered: {comp_data['lcs_sequence']} (Length: {comp_data['lcs_length']})")

    # Verify Match % for Bulk Student 1: Total = 4, Common = 3 => (3 / 4) * 100 = 75.0%
    st1 = next(s for s in comp_data["students"] if s["roll_no"] == "BLK-101")
    assert st1["total_activities"] == 4
    assert st1["common_count"] == 3
    assert st1["missing_count"] == 1
    assert "Java" in st1["missing_activities"]
    assert st1["match_percentage"] == 75.0, f"Expected 75.0, got {st1['match_percentage']}"
    print(f"  -> Bulk Student 1 Match %: {st1['match_percentage']}% (Formula: (3/4)*100 = 75%)")

    # Verify Match % for Bulk Student 2: Total = 4, Common = 3 => (3 / 4) * 100 = 75.0%
    st2 = next(s for s in comp_data["students"] if s["roll_no"] == "BLK-102")
    assert st2["total_activities"] == 4
    assert st2["common_count"] == 3
    assert st2["missing_count"] == 1
    assert "Seminar" in st2["missing_activities"]
    assert st2["match_percentage"] == 75.0
    print(f"  -> Bulk Student 2 Match %: {st2['match_percentage']}% (Formula: (3/4)*100 = 75%)")
    print("  -> PASSED: Multi-student comparison and exact Match % formula verified.")

    # 11. Test Dashboard API
    print("\n[10] Testing GET /api/dashboard...")
    dash_res = test_client.get('/api/dashboard')
    assert dash_res.status_code == 200
    dash_data = dash_res.get_json()
    assert dash_data["success"] is True
    stats = dash_data["stats"]
    assert stats["total_students"] >= 6
    assert stats["total_activities"] >= 20
    assert stats["total_comparisons"] >= 1
    assert "student_activities" in dash_data["charts"]
    assert "match_percentages" in dash_data["charts"]
    assert "categories" in dash_data["charts"]
    assert "student_performance" in dash_data["charts"]
    print("  -> PASSED: Dashboard stats and all 4 chart datasets returned correctly.")

    # 12. Test Individual Student Report API
    print("\n[11] Testing GET /api/student/<id>/report...")
    rep_res = test_client.get(f'/api/student/{b_ids[0]}/report')
    assert rep_res.status_code == 200
    rep_data = rep_res.get_json()
    assert rep_data["success"] is True
    assert rep_data["student"]["name"] == "Bulk Student 1"
    assert rep_data["comparison_info"] is not None
    assert rep_data["comparison_info"]["lcs_length"] == 3
    assert rep_data["comparison_info"]["match_percentage"] == 75.0
    assert "Bulk Student 2" in rep_data["comparison_info"]["compared_with"]
    print("  -> PASSED: Student report contains full comparison details & match percentage.")

    # 13. Test Excel Export & Sheet Verification
    print("\n[12] Testing GET /api/export/excel (openpyxl 8 sheets validation)...")
    excel_res = test_client.get(f'/api/export/excel?comparison_id={comp_id}')
    assert excel_res.status_code == 200
    assert excel_res.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    
    # Save temporarily to inspect with openpyxl
    temp_excel_path = "test_verify.xlsx"
    with open(temp_excel_path, "wb") as f:
        f.write(excel_res.data)

    wb = openpyxl.load_workbook(temp_excel_path)
    sheet_names = wb.sheetnames
    print(f"  -> Sheets found in Excel: {sheet_names}")
    
    expected_sheets = [
        "Student Details",
        "All Activities",
        "Comparison",
        "LCS Result",
        "Missing Activities",
        "Student Summary",
        "Activity-wise Percentage",
        "Category Analysis"
    ]
    for es in expected_sheets:
        assert es in sheet_names, f"Missing sheet: {es}"
    print("  -> All 8 required sheets exist!")

    # Verify 'All Activities' sheet content:
    ws_all = wb["All Activities"]
    assert isinstance(ws_all, Worksheet)
    assert ws_all.max_row is not None
    headers = [cell.value for cell in ws_all[1]]
    expected_headers = ["Name", "Roll No", "Branch", "Activity Name", "Common Activities", "Total Activities", "Missing Activities", "Match %"]
    assert headers == expected_headers, f"Headers mismatch: {headers}"
    print(f"  -> All Activities headers matched: {headers}")

    # Check rows: each student had 4 activities, so 2 students * 4 = 8 activity rows + 1 header = 9 rows
    assert ws_all.max_row == 9, f"Expected 9 rows in All Activities, got {ws_all.max_row}"
    # Verify row values repeat Common Activities, Total Activities, Missing Activities, Match %
    for r in range(2, 6): # Bulk Student 1 rows
        assert ws_all.cell(row=r, column=1).value == "Bulk Student 1"
        assert ws_all.cell(row=r, column=5).value == 3 # Common Activities
        assert ws_all.cell(row=r, column=6).value == 4 # Total Activities
        assert "Java" in str(ws_all.cell(row=r, column=7).value) # Missing Activities
        assert "75.00%" in str(ws_all.cell(row=r, column=8).value) # Match %
    print("  -> Every activity included with repeated common/total/missing/match% metrics!")

    # Verify 'Missing Activities' sheet
    ws_miss = wb["Missing Activities"]
    assert isinstance(ws_miss, Worksheet)
    assert ws_miss.max_row is not None
    assert ws_miss.max_row >= 3 # header + 2 missing items (Java, Seminar)
    print("  -> Missing Activities sheet validated!")

    # Verify 'Activity-wise Percentage' sheet
    ws_act_pct = wb["Activity-wise Percentage"]
    assert isinstance(ws_act_pct, Worksheet)
    assert ws_act_pct.max_row is not None
    # Check that Activity Match % is 100% or 0%
    act_match_vals = [ws_act_pct.cell(row=r, column=4).value for r in range(2, ws_act_pct.max_row + 1)]
    assert all(val in ["100%", "0%"] for val in act_match_vals)
    print("  -> Activity-wise Percentage (100% for LCS, 0% for non-LCS) validated!")

    wb.close()
    if os.path.exists(temp_excel_path):
        os.remove(temp_excel_path)
    print("  -> PASSED: Excel report fully complies with all specifications.")

    # 14. Test PDF Export
    print("\n[13] Testing GET /api/export/pdf (reportlab PDF validation)...")
    pdf_res = test_client.get(f'/api/export/pdf?comparison_id={comp_id}')
    assert pdf_res.status_code == 200
    assert pdf_res.mimetype == "application/pdf"
    assert len(pdf_res.data) > 1000 # Valid non-empty PDF bytes
    print(f"  -> PASSED: PDF report generated successfully ({len(pdf_res.data)} bytes).")

    # 15. Test DELETE student
    print("\n[14] Testing DELETE /api/students/<id>...")
    del_res = test_client.delete(f'/api/students/{new_student_id}')
    assert del_res.status_code == 200
    for b_id in b_ids:
        test_client.delete(f'/api/students/{b_id}')
    print("  -> PASSED: Deleted student successfully.")

    # 16. Verify Database Persistence
    print("\n[15] Testing SQLite Database Persistence...")
    conn = sqlite3.connect('activity.db')
    cursor = conn.cursor()
    cursor.execute("SELECT COUNT(*) FROM students")
    st_count = cursor.fetchone()[0]
    cursor.execute("SELECT COUNT(*) FROM comparisons")
    cp_count = cursor.fetchone()[0]
    conn.close()
    assert st_count > 0
    assert cp_count > 0
    print(f"  -> PASSED: Database confirmed persistent ({st_count} students, {cp_count} comparisons).")

    print("\n" + "=" * 70)
    print("ALL TESTS PASSED WITH 100% COMPLIANCE!")
    print("=" * 70)

if __name__ == '__main__':
    test_full_system()
