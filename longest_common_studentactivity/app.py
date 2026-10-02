import os
import json
import sqlite3
import datetime
from typing import Any
from flask import Flask, render_template, request, jsonify, send_file
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ReportLab imports for PDF generation
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas

# Base directories
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, 'activity.db')
REPORTS_PDF_DIR = os.path.join(BASE_DIR, 'reports', 'pdf')
REPORTS_EXCEL_DIR = os.path.join(BASE_DIR, 'reports', 'excel')

os.makedirs(REPORTS_PDF_DIR, exist_ok=True)
os.makedirs(REPORTS_EXCEL_DIR, exist_ok=True)

app = Flask(__name__, template_folder='templates', static_folder='static')


# ==============================================================================
# DATABASE SETUP & HELPERS
# ==============================================================================

def get_db_connection():
    """Returns a SQLite connection with foreign keys enabled and row factory set to Row."""
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON;")
    return conn

def init_db():
    """Initializes SQLite database and creates all required tables if they don't exist."""
    conn = get_db_connection()
    cursor = conn.cursor()

    # 1. students table
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS students (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            roll_no TEXT UNIQUE NOT NULL,
            branch TEXT NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    ''')

    # 2. activities table
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS activities (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            student_id INTEGER NOT NULL,
            activity_name TEXT NOT NULL,
            category TEXT NOT NULL,
            activity_date TEXT NOT NULL,
            verification_status TEXT NOT NULL,
            proof_reference TEXT,
            order_index INTEGER DEFAULT 0,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (student_id) REFERENCES students (id) ON DELETE CASCADE
        )
    ''')

    # 3. comparisons table
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS comparisons (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            comparison_date TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            student_ids TEXT NOT NULL,
            lcs_sequence TEXT NOT NULL,
            lcs_length INTEGER NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    ''')

    # 4. comparison_students table
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS comparison_students (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            comparison_id INTEGER NOT NULL,
            student_id INTEGER NOT NULL,
            student_name TEXT NOT NULL,
            roll_no TEXT NOT NULL,
            total_activities INTEGER NOT NULL,
            common_activities_count INTEGER NOT NULL,
            missing_activities_list TEXT NOT NULL,
            match_percentage REAL NOT NULL,
            FOREIGN KEY (comparison_id) REFERENCES comparisons (id) ON DELETE CASCADE
        )
    ''')

    # 5. comparison_results table
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS comparison_results (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            comparison_id INTEGER NOT NULL,
            student_id INTEGER,
            common_activity TEXT NOT NULL,
            lcs_position INTEGER NOT NULL,
            FOREIGN KEY (comparison_id) REFERENCES comparisons (id) ON DELETE CASCADE
        )
    ''')

    # 6. missing_activities table
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS missing_activities (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            comparison_id INTEGER NOT NULL,
            student_id INTEGER NOT NULL,
            student_name TEXT NOT NULL,
            roll_no TEXT NOT NULL,
            missing_activity TEXT NOT NULL,
            FOREIGN KEY (comparison_id) REFERENCES comparisons (id) ON DELETE CASCADE
        )
    ''')

    # 7. comparison_activity_results table
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS comparison_activity_results (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            comparison_id INTEGER NOT NULL,
            student_id INTEGER NOT NULL,
            activity_name TEXT NOT NULL,
            is_common INTEGER NOT NULL,
            activity_match_pct REAL NOT NULL,
            student_match_pct REAL NOT NULL,
            FOREIGN KEY (comparison_id) REFERENCES comparisons (id) ON DELETE CASCADE
        )
    ''')

    conn.commit()

    # Seed initial data if students table is empty
    cursor.execute("SELECT COUNT(*) FROM students")
    count = cursor.fetchone()[0]
    if count == 0:
        seed_sample_data(cursor, conn)
    else:
        conn.close()

def seed_sample_data(cursor, conn):
    """Seeds rich default students and ordered activities for instant demonstration."""
    sample_students = [
        {
            "name": "Trisha Sharma",
            "roll_no": "101",
            "branch": "Computer Science & Engineering",
            "activities": [
                ("Python Programming", "Coding", "2026-08-10", "Verified", "CERT-PY-8831"),
                ("Data Structures & Algorithms", "Academic", "2026-08-20", "Verified", "DSA-LAB-402"),
                ("Smart India Hackathon", "Hackathon", "2026-09-02", "Verified", "SIH-TEAM-01"),
                ("Java Enterprise", "Coding", "2026-09-12", "Verified", "CERT-JV-992"),
                ("AI Seminar", "Seminar", "2026-09-18", "Verified", "SEM-REF-512"),
                ("Cloud Computing Workshop", "Workshop", "2026-09-25", "Pending", "AWS-ATTEND-12")
            ]
        },
        {
            "name": "Rahul Verma",
            "roll_no": "102",
            "branch": "Information Technology",
            "activities": [
                ("Python Programming", "Coding", "2026-08-12", "Verified", "CERT-PY-9011"),
                ("Data Structures & Algorithms", "Academic", "2026-08-22", "Verified", "DSA-LAB-411"),
                ("Smart India Hackathon", "Hackathon", "2026-09-03", "Verified", "SIH-TEAM-02"),
                ("Java Enterprise", "Coding", "2026-09-15", "Verified", "CERT-JV-102"),
                ("College Sports Tournament", "Sports", "2026-09-21", "Verified", "TROPHY-SPT-09"),
                ("Cybersecurity Workshop", "Workshop", "2026-09-27", "Pending", "CYBER-WS-33")
            ]
        },
        {
            "name": "Priya Nair",
            "roll_no": "103",
            "branch": "Artificial Intelligence & Data Science",
            "activities": [
                ("Python Programming", "Coding", "2026-08-11", "Verified", "PY-AI-441"),
                ("Data Structures & Algorithms", "Academic", "2026-08-21", "Verified", "DSA-LAB-408"),
                ("Smart India Hackathon", "Hackathon", "2026-09-04", "Verified", "SIH-TEAM-04"),
                ("Web Development Bootcamp", "Certification", "2026-09-14", "Verified", "UDEMY-WD-88"),
                ("Machine Learning Project", "Project", "2026-09-24", "Verified", "GITHUB-ML-REPO"),
                ("AI Seminar", "Seminar", "2026-09-28", "Pending", "SEM-REF-599")
            ]
        },
        {
            "name": "Amit Patel",
            "roll_no": "104",
            "branch": "Electronics & Communication",
            "activities": [
                ("Python Programming", "Coding", "2026-08-15", "Verified", "CERT-PY-552"),
                ("Embedded C Systems", "Academic", "2026-08-25", "Verified", "EMB-LAB-201"),
                ("Data Structures & Algorithms", "Academic", "2026-09-01", "Verified", "DSA-LAB-419"),
                ("Smart India Hackathon", "Hackathon", "2026-09-05", "Verified", "SIH-TEAM-06"),
                ("Robotics Championship", "Competition", "2026-09-20", "Verified", "ROBO-MEDAL-04")
            ]
        }
    ]

    for st in sample_students:
        cursor.execute(
            "INSERT INTO students (name, roll_no, branch) VALUES (?, ?, ?)",
            (st["name"], st["roll_no"], st["branch"])
        )
        student_id = cursor.lastrowid
        for idx, act in enumerate(st["activities"]):
            cursor.execute(
                """INSERT INTO activities 
                   (student_id, activity_name, category, activity_date, verification_status, proof_reference, order_index)
                   VALUES (?, ?, ?, ?, ?, ?, ?)""",
                (student_id, act[0], act[1], act[2], act[3], act[4], idx)
            )

    conn.commit()
    conn.close()

# Initialize Database on app start
init_db()

# ==============================================================================
# DYNAMIC PROGRAMMING LCS ALGORITHM & COMPARISON LOGIC
# ==============================================================================

def lcs_two_sequences(seq1, seq2):
    """
    Computes Longest Common Subsequence of two sequences using 2D Dynamic Programming.
    Returns the ordered list of items in the LCS.
    """
    m, n = len(seq1), len(seq2)
    dp = [[0] * (n + 1) for _ in range(m + 1)]

    for i in range(1, m + 1):
        for j in range(1, n + 1):
            if seq1[i - 1].strip().lower() == seq2[j - 1].strip().lower():
                dp[i][j] = dp[i - 1][j - 1] + 1
            else:
                dp[i][j] = max(dp[i - 1][j], dp[i][j - 1])

    # Backtracking to reconstruct LCS sequence
    lcs = []
    i, j = m, n
    while i > 0 and j > 0:
        if seq1[i - 1].strip().lower() == seq2[j - 1].strip().lower():
            lcs.append(seq1[i - 1].strip())
            i -= 1
            j -= 1
        elif dp[i - 1][j] >= dp[i][j - 1]:
            i -= 1
        else:
            j -= 1

    lcs.reverse()
    return lcs

def multi_student_lcs(sequences_list):
    """
    Computes common sequence across K student sequences using progressive DP LCS.
    Guarantees common ordering across all selected students.
    """
    if not sequences_list:
        return []
    current_lcs = [act.strip() for act in sequences_list[0]]
    for seq in sequences_list[1:]:
        current_lcs = lcs_two_sequences(current_lcs, seq)
        if not current_lcs:
            break
    return current_lcs

def compute_student_lcs_stats(student_activities, lcs_items):
    """
    Aligns student's activities against common LCS sequence.
    Identifies matched common activities and missing activities.
    
    EXACT MATCH % FORMULA (Section 13):
    Match % = (Common Activities / Total Activities of that student) * 100
    """
    total_count = len(student_activities)
    matched_indices = set()
    lcs_idx = 0
    num_lcs = len(lcs_items)

    for idx, act in enumerate(student_activities):
        if lcs_idx < num_lcs and act.strip().lower() == lcs_items[lcs_idx].strip().lower():
            matched_indices.add(idx)
            lcs_idx += 1

    common_activities = [student_activities[idx] for idx in sorted(matched_indices)]
    missing_activities = [student_activities[idx] for idx in range(total_count) if idx not in matched_indices]

    common_count = len(common_activities)
    # Strictly apply formula
    match_pct = (common_count / total_count * 100.0) if total_count > 0 else 0.0
    match_pct = round(match_pct, 2)

    return {
        "total_activities": total_count,
        "common_activities": common_activities,
        "common_count": common_count,
        "missing_activities": missing_activities,
        "missing_count": len(missing_activities),
        "match_percentage": match_pct,
        "matched_indices": sorted(list(matched_indices))
    }

def execute_comparison(student_ids):
    """
    Performs multi-student comparison, computes LCS, stores results in SQLite,
    and returns full comparison dictionary.
    """
    if not student_ids or len(student_ids) < 2:
        raise ValueError("At least 2 students must be selected for comparison.")

    conn = get_db_connection()
    cursor = conn.cursor()

    # Fetch students
    placeholders = ','.join(['?'] * len(student_ids))
    cursor.execute(f"SELECT * FROM students WHERE id IN ({placeholders})", student_ids)
    students_db = cursor.fetchall()
    student_map = {s['id']: dict(s) for s in students_db}

    # Verify all students found
    valid_ids = [s_id for s_id in student_ids if s_id in student_map]
    if len(valid_ids) < 2:
        conn.close()
        raise ValueError("Could not find at least 2 valid students for comparison.")

    # Fetch activities for each student in sequence order
    students_data = []
    sequences_list = []

    for s_id in valid_ids:
        s_info = student_map[s_id]
        cursor.execute(
            "SELECT * FROM activities WHERE student_id = ? ORDER BY order_index ASC, id ASC",
            (s_id,)
        )
        acts = cursor.fetchall()
        act_names = [a['activity_name'].strip() for a in acts]
        students_data.append({
            "student_id": s_id,
            "name": s_info['name'],
            "roll_no": s_info['roll_no'],
            "branch": s_info['branch'],
            "raw_activities": [dict(a) for a in acts],
            "activity_names": act_names
        })
        sequences_list.append(act_names)

    # Compute DP Multi-Student LCS
    lcs_sequence = multi_student_lcs(sequences_list)
    lcs_length = len(lcs_sequence)
    lcs_sequence_str = json.dumps(lcs_sequence)
    student_ids_str = json.dumps(valid_ids)

    # Insert into comparisons table
    cursor.execute(
        "INSERT INTO comparisons (student_ids, lcs_sequence, lcs_length) VALUES (?, ?, ?)",
        (student_ids_str, lcs_sequence_str, lcs_length)
    )
    comparison_id = cursor.lastrowid

    # For each common activity in LCS, record in comparison_results
    for pos, common_act in enumerate(lcs_sequence):
        cursor.execute(
            """INSERT INTO comparison_results (comparison_id, student_id, common_activity, lcs_position)
               VALUES (?, NULL, ?, ?)""",
            (comparison_id, common_act, pos + 1)
        )

    # Analyze each student
    processed_students = []
    for s in students_data:
        stats = compute_student_lcs_stats(s["activity_names"], lcs_sequence)
        matched_set = set(stats["matched_indices"])
        missing_str = ", ".join(stats["missing_activities"]) if stats["missing_activities"] else "None"

        # Record in comparison_students
        cursor.execute(
            """INSERT INTO comparison_students 
               (comparison_id, student_id, student_name, roll_no, total_activities, common_activities_count, missing_activities_list, match_percentage)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                comparison_id,
                s["student_id"],
                s["name"],
                s["roll_no"],
                stats["total_activities"],
                stats["common_count"],
                missing_str,
                stats["match_percentage"]
            )
        )

        # Record missing activities row-by-row
        for miss_act in stats["missing_activities"]:
            cursor.execute(
                """INSERT INTO missing_activities (comparison_id, student_id, student_name, roll_no, missing_activity)
                   VALUES (?, ?, ?, ?, ?)""",
                (comparison_id, s["student_id"], s["name"], s["roll_no"], miss_act)
            )

        # Record activity-wise percentage
        for idx, act_obj in enumerate(s["raw_activities"]):
            act_name = act_obj["activity_name"]
            is_common = 1 if idx in matched_set else 0
            act_match_pct = 100.0 if is_common else 0.0
            cursor.execute(
                """INSERT INTO comparison_activity_results 
                   (comparison_id, student_id, activity_name, is_common, activity_match_pct, student_match_pct)
                   VALUES (?, ?, ?, ?, ?, ?)""",
                (comparison_id, s["student_id"], act_name, is_common, act_match_pct, stats["match_percentage"])
            )

        processed_students.append({
            "student_id": s["student_id"],
            "name": s["name"],
            "roll_no": s["roll_no"],
            "branch": s["branch"],
            "total_activities": stats["total_activities"],
            "common_activities": stats["common_activities"],
            "common_count": stats["common_count"],
            "missing_activities": stats["missing_activities"],
            "missing_count": stats["missing_count"],
            "match_percentage": stats["match_percentage"],
            "activities": s["raw_activities"]
        })

    conn.commit()
    conn.close()

    return {
        "comparison_id": comparison_id,
        "comparison_date": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "lcs_sequence": lcs_sequence,
        "lcs_length": lcs_length,
        "common_activities": lcs_sequence,
        "students": processed_students
    }

# ==============================================================================
# FLASK ROUTE: SINGLE PAGE APPLICATION
# ==============================================================================

@app.route('/')
def index():
    """Renders the SINGLE page of the application. Everything lives here."""
    return render_template('index.html')

# ==============================================================================
# INTERNAL REST APIS: STUDENTS CRUD
# ==============================================================================

@app.route('/api/students', methods=['GET'])
def get_students():
    """Returns list of all students with activity counts and latest match %."""
    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute("""
            SELECT s.*, 
                   COUNT(a.id) AS total_activities
            FROM students s
            LEFT JOIN activities a ON s.id = a.student_id
            GROUP BY s.id
            ORDER BY s.id ASC
        """)
        students_raw = [dict(row) for row in cursor.fetchall()]

        # Attach latest match percentage from comparison_students
        for st in students_raw:
            cursor.execute("""
                SELECT match_percentage, comparison_id, created_at 
                FROM comparison_students cs
                JOIN comparisons c ON cs.comparison_id = c.id
                WHERE cs.student_id = ?
                ORDER BY c.id DESC LIMIT 1
            """, (st['id'],))
            latest_comp = cursor.fetchone()
            st['latest_match_pct'] = latest_comp['match_percentage'] if latest_comp else 0.0

        conn.close()
        return jsonify({"success": True, "students": students_raw})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500

@app.route('/api/students', methods=['POST'])
def add_student():
    """Adds a new student."""
    try:
        data = request.get_json() or {}
        name = str(data.get('name', '')).strip()
        roll_no = str(data.get('roll_no', '')).strip()
        branch = str(data.get('branch', '')).strip()

        if not name:
            return jsonify({"success": False, "error": "Student name cannot be empty."}), 400
        if not roll_no:
            return jsonify({"success": False, "error": "Roll number cannot be empty."}), 400
        if not branch:
            return jsonify({"success": False, "error": "Branch cannot be empty."}), 400

        conn = get_db_connection()
        cursor = conn.cursor()

        # Check duplicate roll number
        cursor.execute("SELECT id FROM students WHERE roll_no = ?", (roll_no,))
        if cursor.fetchone():
            conn.close()
            return jsonify({"success": False, "error": f"A student with roll number '{roll_no}' already exists."}), 400

        cursor.execute(
            "INSERT INTO students (name, roll_no, branch) VALUES (?, ?, ?)",
            (name, roll_no, branch)
        )
        student_id = cursor.lastrowid
        conn.commit()
        conn.close()

        return jsonify({
            "success": True,
            "message": f"Student '{name}' added successfully.",
            "student_id": student_id
        }), 201
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500

@app.route('/api/students/bulk', methods=['POST'])
def add_students_bulk():
    """Dynamic bulk student creation from first screen."""
    try:
        data = request.get_json() or {}
        students_list = data.get('students', [])

        if not students_list or not isinstance(students_list, list):
            return jsonify({"success": False, "error": "Invalid student list provided."}), 400

        conn = get_db_connection()
        cursor = conn.cursor()

        created_ids = []
        for idx, s in enumerate(students_list):
            name = str(s.get('name', '')).strip()
            roll_no = str(s.get('roll_no', '')).strip()
            branch = str(s.get('branch', '')).strip()
            activities = s.get('activities', [])

            if not name or not roll_no or not branch:
                conn.close()
                return jsonify({"success": False, "error": f"Student #{idx+1} has missing fields (Name, Roll No, Branch are required)."}), 400

            # Check duplicate roll_no
            cursor.execute("SELECT id FROM students WHERE roll_no = ?", (roll_no,))
            if cursor.fetchone():
                conn.close()
                return jsonify({"success": False, "error": f"Roll number '{roll_no}' already exists."}), 400

            cursor.execute(
                "INSERT INTO students (name, roll_no, branch) VALUES (?, ?, ?)",
                (name, roll_no, branch)
            )
            student_id = cursor.lastrowid
            created_ids.append(student_id)

            # Insert any initial activities provided
            for a_idx, act in enumerate(activities):
                act_name = str(act.get('activity_name', '')).strip()
                cat = str(act.get('category', 'Coding')).strip() or 'Coding'
                dt = str(act.get('activity_date', datetime.date.today().strftime('%Y-%m-%d'))).strip()
                st_val = str(act.get('verification_status', 'Verified')).strip() or 'Verified'
                proof = str(act.get('proof_reference', '')).strip()

                if act_name:
                    cursor.execute("""
                        INSERT INTO activities 
                        (student_id, activity_name, category, activity_date, verification_status, proof_reference, order_index)
                        VALUES (?, ?, ?, ?, ?, ?, ?)
                    """, (student_id, act_name, cat, dt, st_val, proof, a_idx))

        conn.commit()
        conn.close()

        return jsonify({
            "success": True,
            "message": f"Successfully created {len(created_ids)} students.",
            "student_ids": created_ids
        }), 201
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500

@app.route('/api/students/<int:student_id>', methods=['PUT'])
def update_student(student_id):
    """Updates an existing student."""
    try:
        data = request.get_json() or {}
        name = str(data.get('name', '')).strip()
        roll_no = str(data.get('roll_no', '')).strip()
        branch = str(data.get('branch', '')).strip()

        if not name or not roll_no or not branch:
            return jsonify({"success": False, "error": "Name, roll number, and branch are all required."}), 400

        conn = get_db_connection()
        cursor = conn.cursor()

        # Check student exists
        cursor.execute("SELECT id FROM students WHERE id = ?", (student_id,))
        if not cursor.fetchone():
            conn.close()
            return jsonify({"success": False, "error": "Student not found."}), 404

        # Check duplicate roll_no for other students
        cursor.execute("SELECT id FROM students WHERE roll_no = ? AND id != ?", (roll_no, student_id))
        if cursor.fetchone():
            conn.close()
            return jsonify({"success": False, "error": f"Another student already has roll number '{roll_no}'."}), 400

        cursor.execute(
            "UPDATE students SET name = ?, roll_no = ?, branch = ? WHERE id = ?",
            (name, roll_no, branch, student_id)
        )
        conn.commit()
        conn.close()

        return jsonify({"success": True, "message": "Student updated successfully."})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500

@app.route('/api/students/<int:student_id>', methods=['DELETE'])
def delete_student(student_id):
    """Deletes a student and their activities (cascade)."""
    try:
        conn = get_db_connection()
        cursor = conn.cursor()

        cursor.execute("SELECT name FROM students WHERE id = ?", (student_id,))
        row = cursor.fetchone()
        if not row:
            conn.close()
            return jsonify({"success": False, "error": "Student not found."}), 404

        name = row['name']
        cursor.execute("DELETE FROM students WHERE id = ?", (student_id,))
        conn.commit()
        conn.close()

        return jsonify({"success": True, "message": f"Student '{name}' deleted successfully."})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500

# ==============================================================================
# INTERNAL REST APIS: ACTIVITIES CRUD & HISTORY
# ==============================================================================

@app.route('/api/activities', methods=['GET'])
def get_activities():
    """Lists activities with optional filters by student_id, category, search, and date."""
    try:
        student_id = request.args.get('student_id')
        category = request.args.get('category')
        search = request.args.get('search', '').strip()
        start_date = request.args.get('start_date')
        end_date = request.args.get('end_date')

        conn = get_db_connection()
        cursor = conn.cursor()

        query = """
            SELECT a.*, s.name AS student_name, s.roll_no, s.branch
            FROM activities a
            JOIN students s ON a.student_id = s.id
            WHERE 1=1
        """
        params = []

        if student_id:
            query += " AND a.student_id = ?"
            params.append(student_id)
        if category and category != 'All':
            query += " AND a.category = ?"
            params.append(category)
        if search:
            query += " AND (a.activity_name LIKE ? OR s.name LIKE ? OR s.roll_no LIKE ?)"
            term = f"%{search}%"
            params.extend([term, term, term])
        if start_date:
            query += " AND a.activity_date >= ?"
            params.append(start_date)
        if end_date:
            query += " AND a.activity_date <= ?"
            params.append(end_date)

        query += " ORDER BY a.student_id ASC, a.order_index ASC, a.id ASC"

        cursor.execute(query, params)
        acts = [dict(row) for row in cursor.fetchall()]
        conn.close()

        return jsonify({"success": True, "activities": acts})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500

@app.route('/api/activities', methods=['POST'])
def add_activity():
    """Adds a new activity for a student."""
    try:
        data = request.get_json() or {}
        student_id = data.get('student_id')
        activity_name = str(data.get('activity_name', '')).strip()
        category = str(data.get('category', '')).strip() or 'Coding'
        activity_date = str(data.get('activity_date', '')).strip() or datetime.date.today().strftime('%Y-%m-%d')
        verification_status = str(data.get('verification_status', '')).strip() or 'Verified'
        proof_reference = str(data.get('proof_reference', '')).strip()

        if not student_id:
            return jsonify({"success": False, "error": "Please select a valid student."}), 400
        if not activity_name:
            return jsonify({"success": False, "error": "Activity name cannot be empty."}), 400

        conn = get_db_connection()
        cursor = conn.cursor()

        # Check student exists
        cursor.execute("SELECT id FROM students WHERE id = ?", (student_id,))
        if not cursor.fetchone():
            conn.close()
            return jsonify({"success": False, "error": "Selected student does not exist."}), 404

        # Compute next order_index
        cursor.execute("SELECT COALESCE(MAX(order_index), -1) + 1 FROM activities WHERE student_id = ?", (student_id,))
        next_order = cursor.fetchone()[0]

        cursor.execute("""
            INSERT INTO activities 
            (student_id, activity_name, category, activity_date, verification_status, proof_reference, order_index)
            VALUES (?, ?, ?, ?, ?, ?, ?)
        """, (student_id, activity_name, category, activity_date, verification_status, proof_reference, next_order))

        act_id = cursor.lastrowid
        conn.commit()
        conn.close()

        return jsonify({
            "success": True,
            "message": "Activity added successfully.",
            "activity_id": act_id
        }), 201
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500

@app.route('/api/activities/<int:activity_id>', methods=['PUT'])
def update_activity(activity_id):
    """Updates an activity."""
    try:
        data = request.get_json() or {}
        activity_name = str(data.get('activity_name', '')).strip()
        category = str(data.get('category', '')).strip()
        activity_date = str(data.get('activity_date', '')).strip()
        verification_status = str(data.get('verification_status', '')).strip()
        proof_reference = str(data.get('proof_reference', '')).strip()

        if not activity_name:
            return jsonify({"success": False, "error": "Activity name cannot be empty."}), 400

        conn = get_db_connection()
        cursor = conn.cursor()

        cursor.execute("SELECT id FROM activities WHERE id = ?", (activity_id,))
        if not cursor.fetchone():
            conn.close()
            return jsonify({"success": False, "error": "Activity not found."}), 404

        cursor.execute("""
            UPDATE activities 
            SET activity_name = ?, category = ?, activity_date = ?, verification_status = ?, proof_reference = ?
            WHERE id = ?
        """, (activity_name, category, activity_date, verification_status, proof_reference, activity_id))

        conn.commit()
        conn.close()

        return jsonify({"success": True, "message": "Activity updated successfully."})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500

@app.route('/api/activities/<int:activity_id>', methods=['DELETE'])
def delete_activity(activity_id):
    """Deletes an activity."""
    try:
        conn = get_db_connection()
        cursor = conn.cursor()

        cursor.execute("SELECT id FROM activities WHERE id = ?", (activity_id,))
        if not cursor.fetchone():
            conn.close()
            return jsonify({"success": False, "error": "Activity not found."}), 404

        cursor.execute("DELETE FROM activities WHERE id = ?", (activity_id,))
        conn.commit()
        conn.close()

        return jsonify({"success": True, "message": "Activity deleted successfully."})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500

# ==============================================================================
# INTERNAL REST APIS: COMPARISON & LCS
# ==============================================================================

@app.route('/api/compare', methods=['POST'])
def compare_students_api():
    """Performs multi-student LCS comparison and returns comprehensive stats."""
    try:
        data = request.get_json() or {}
        student_ids = data.get('student_ids', [])

        if not student_ids or len(student_ids) < 2:
            return jsonify({"success": False, "error": "Please select at least 2 students to compare."}), 400

        result = execute_comparison(student_ids)
        return jsonify({"success": True, "data": result})
    except ValueError as ve:
        return jsonify({"success": False, "error": str(ve)}), 400
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500

@app.route('/api/comparisons', methods=['GET'])
def get_comparisons():
    """Lists comparison history."""
    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute("""
            SELECT c.*, 
                   COUNT(cs.id) as student_count
            FROM comparisons c
            LEFT JOIN comparison_students cs ON c.id = cs.comparison_id
            GROUP BY c.id
            ORDER BY c.id DESC
        """)
        rows = cursor.fetchall()
        comparisons = []
        for r in rows:
            comp_dict = dict(r)
            try:
                comp_dict['lcs_sequence'] = json.loads(r['lcs_sequence'])
                comp_dict['student_ids'] = json.loads(r['student_ids'])
            except Exception:
                pass
            comparisons.append(comp_dict)

        conn.close()
        return jsonify({"success": True, "comparisons": comparisons})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500

@app.route('/api/comparisons/<int:comparison_id>', methods=['GET'])
def get_comparison_detail(comparison_id):
    """Retrieves full details of a specific comparison."""
    try:
        conn = get_db_connection()
        cursor = conn.cursor()

        cursor.execute("SELECT * FROM comparisons WHERE id = ?", (comparison_id,))
        comp = cursor.fetchone()
        if not comp:
            conn.close()
            return jsonify({"success": False, "error": "Comparison not found."}), 404

        comp_dict = dict(comp)
        try:
            comp_dict['lcs_sequence'] = json.loads(comp['lcs_sequence'])
            comp_dict['student_ids'] = json.loads(comp['student_ids'])
        except Exception:
            pass

        # Fetch student details in this comparison
        cursor.execute("""
            SELECT cs.*, s.branch
            FROM comparison_students cs
            JOIN students s ON cs.student_id = s.id
            WHERE cs.comparison_id = ?
        """, (comparison_id,))
        students = [dict(s) for s in cursor.fetchall()]

        for st in students:
            # Parse missing activities list
            miss_str = st.get('missing_activities_list', '')
            st['missing_activities'] = [m.strip() for m in miss_str.split(',') if m.strip() and m.strip() != 'None']
            st['common_activities'] = comp_dict['lcs_sequence']

        # Fetch missing activities table rows
        cursor.execute("SELECT * FROM missing_activities WHERE comparison_id = ?", (comparison_id,))
        missing_rows = [dict(m) for m in cursor.fetchall()]

        conn.close()
        return jsonify({
            "success": True,
            "comparison": comp_dict,
            "students": students,
            "missing_activities": missing_rows
        })
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500

# ==============================================================================
# INTERNAL REST APIS: DASHBOARD & INDIVIDUAL STUDENT REPORT
# ==============================================================================

@app.route('/api/dashboard', methods=['GET'])
def get_dashboard_data():
    """Computes all summary statistics and chart datasets for dashboard."""
    try:
        conn = get_db_connection()
        cursor = conn.cursor()

        # 1. Total Students
        cursor.execute("SELECT COUNT(*) FROM students")
        total_students = cursor.fetchone()[0]

        # 2. Total Activities
        cursor.execute("SELECT COUNT(*) FROM activities")
        total_activities = cursor.fetchone()[0]

        # 3. Total Comparisons
        cursor.execute("SELECT COUNT(*) FROM comparisons")
        total_comparisons = cursor.fetchone()[0]

        # 4. Latest LCS Length
        cursor.execute("SELECT lcs_length FROM comparisons ORDER BY id DESC LIMIT 1")
        latest_lcs_row = cursor.fetchone()
        latest_lcs_length = latest_lcs_row[0] if latest_lcs_row else 0

        # 5. Average Match % from latest comparison or all comparison students
        cursor.execute("SELECT AVG(match_percentage) FROM comparison_students")
        avg_match_row = cursor.fetchone()
        avg_match_pct = round(avg_match_row[0], 2) if (avg_match_row and avg_match_row[0] is not None) else 0.0

        # Chart 1: Student Activity Graph (Student Name vs Total Activities)
        cursor.execute("""
            SELECT s.name, COUNT(a.id) as act_count
            FROM students s
            LEFT JOIN activities a ON s.id = a.student_id
            GROUP BY s.id
            ORDER BY act_count DESC, s.name ASC
        """)
        student_act_chart = [{"name": r[0], "count": r[1]} for r in cursor.fetchall()]

        # Chart 2: Match Percentage Graph (Student Name vs Match %)
        cursor.execute("""
            SELECT s.name, COALESCE(
                (SELECT cs.match_percentage 
                 FROM comparison_students cs 
                 JOIN comparisons c ON cs.comparison_id = c.id 
                 WHERE cs.student_id = s.id 
                 ORDER BY c.id DESC LIMIT 1), 0.0
            ) as match_pct
            FROM students s
            ORDER BY s.name ASC
        """)
        match_pct_chart = [{"name": r[0], "match_pct": round(r[1], 2)} for r in cursor.fetchall()]

        # Chart 3: Activity Category Graph (Category vs Number of Activities)
        cursor.execute("""
            SELECT category, COUNT(*) as cat_count
            FROM activities
            GROUP BY category
            ORDER BY cat_count DESC
        """)
        category_chart = [{"category": r[0], "count": r[1]} for r in cursor.fetchall()]

        # Chart 4: Individual Student Performance (Total, Common, Missing, Match %)
        # Fetch from latest comparison for each student
        cursor.execute("""
            SELECT s.name, 
                   COUNT(a.id) AS total_activities,
                   COALESCE((SELECT cs.common_activities_count FROM comparison_students cs JOIN comparisons c ON cs.comparison_id = c.id WHERE cs.student_id = s.id ORDER BY c.id DESC LIMIT 1), 0) AS common_activities,
                   COALESCE((SELECT cs.match_percentage FROM comparison_students cs JOIN comparisons c ON cs.comparison_id = c.id WHERE cs.student_id = s.id ORDER BY c.id DESC LIMIT 1), 0.0) AS match_percentage
            FROM students s
            LEFT JOIN activities a ON s.id = a.student_id
            GROUP BY s.id
            ORDER BY s.name ASC
        """)
        perf_rows = cursor.fetchall()
        student_perf_chart = []
        for r in perf_rows:
            tot = r[1]
            comm = r[2]
            miss = max(0, tot - comm)
            mp = round(r[3], 2)
            student_perf_chart.append({
                "name": r[0],
                "total": tot,
                "common": comm,
                "missing": miss,
                "match_pct": mp
            })

        conn.close()

        return jsonify({
            "success": True,
            "stats": {
                "total_students": total_students,
                "total_activities": total_activities,
                "total_comparisons": total_comparisons,
                "latest_lcs_length": latest_lcs_length,
                "average_match_pct": avg_match_pct
            },
            "charts": {
                "student_activities": student_act_chart,
                "match_percentages": match_pct_chart,
                "categories": category_chart,
                "student_performance": student_perf_chart
            }
        })
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500

@app.route('/api/student/<int:student_id>/report', methods=['GET'])
def get_student_report(student_id):
    """
    Returns full individual student report including comparison details,
    LCS sequence, common activities, missing activities, and exact Match %.
    """
    try:
        conn = get_db_connection()
        cursor = conn.cursor()

        # Student info
        cursor.execute("SELECT * FROM students WHERE id = ?", (student_id,))
        student = cursor.fetchone()
        if not student:
            conn.close()
            return jsonify({"success": False, "error": "Student not found."}), 404

        student_dict = dict(student)

        # All activities of this student
        cursor.execute("SELECT * FROM activities WHERE student_id = ? ORDER BY order_index ASC, id ASC", (student_id,))
        activities = [dict(a) for a in cursor.fetchall()]
        student_dict['activities'] = activities
        student_dict['total_activities'] = len(activities)

        # Latest comparison involving this student
        cursor.execute("""
            SELECT c.*, cs.common_activities_count, cs.missing_activities_list, cs.match_percentage
            FROM comparison_students cs
            JOIN comparisons c ON cs.comparison_id = c.id
            WHERE cs.student_id = ?
            ORDER BY c.id DESC LIMIT 1
        """, (student_id,))
        latest_comp = cursor.fetchone()

        comparison_info = None
        if latest_comp:
            comp_id = latest_comp['id']
            try:
                lcs_seq = json.loads(latest_comp['lcs_sequence'])
                all_student_ids = json.loads(latest_comp['student_ids'])
            except Exception:
                lcs_seq = []
                all_student_ids = []

            # Find compared with student names
            other_ids = [sid for sid in all_student_ids if sid != student_id]
            other_names = []
            if other_ids:
                ph = ','.join(['?'] * len(other_ids))
                cursor.execute(f"SELECT name FROM students WHERE id IN ({ph})", other_ids)
                other_names = [row['name'] for row in cursor.fetchall()]

            # Missing activities list
            miss_str = latest_comp['missing_activities_list']
            missing_list = [m.strip() for m in miss_str.split(',') if m.strip() and m.strip() != 'None']

            # Match % strictly using exact formula
            common_count = latest_comp['common_activities_count']
            tot_act = len(activities)
            calc_match_pct = (common_count / tot_act * 100.0) if tot_act > 0 else 0.0

            comparison_info = {
                "comparison_id": comp_id,
                "comparison_date": latest_comp['comparison_date'],
                "compared_with": other_names,
                "common_activity_sequence": lcs_seq,
                "lcs_length": latest_comp['lcs_length'],
                "total_activities": tot_act,
                "common_activities_count": common_count,
                "common_activities": lcs_seq,
                "missing_activities": missing_list,
                "match_percentage": round(calc_match_pct, 2)
            }

        # Full comparison history for this student
        cursor.execute("""
            SELECT c.id AS comparison_id, c.comparison_date, c.lcs_length, c.lcs_sequence,
                   cs.common_activities_count, cs.missing_activities_list, cs.match_percentage
            FROM comparison_students cs
            JOIN comparisons c ON cs.comparison_id = c.id
            WHERE cs.student_id = ?
            ORDER BY c.id DESC
        """, (student_id,))
        history_rows = cursor.fetchall()
        comp_history = []
        for h in history_rows:
            h_dict = dict(h)
            try:
                h_dict['lcs_sequence'] = json.loads(h['lcs_sequence'])
            except Exception:
                pass
            comp_history.append(h_dict)

        conn.close()

        return jsonify({
            "success": True,
            "student": student_dict,
            "comparison_info": comparison_info,
            "comparison_history": comp_history
        })
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500

# ==============================================================================
# EXCEL REPORT GENERATION (openpyxl) - SECTIONS 17, 18, 19, 20
# ==============================================================================

def get_or_run_active_comparison(comparison_id=None):
    """Fetches comparison data for export. If none specified, takes latest or runs on all students."""
    conn = get_db_connection()
    cursor = conn.cursor()

    if comparison_id:
        cursor.execute("SELECT id FROM comparisons WHERE id = ?", (comparison_id,))
        row = cursor.fetchone()
        if not row:
            conn.close()
            raise ValueError(f"Comparison with ID {comparison_id} not found.")
        target_id = comparison_id
    else:
        cursor.execute("SELECT id FROM comparisons ORDER BY id DESC LIMIT 1")
        row = cursor.fetchone()
        if row:
            target_id = row['id']
        else:
            # No comparisons exist yet. Run comparison on all students!
            cursor.execute("SELECT id FROM students ORDER BY id ASC")
            all_sids = [r['id'] for r in cursor.fetchall()]
            conn.close()
            if len(all_sids) < 2:
                raise ValueError("At least 2 students must exist in database to create a comparison report.")
            res = execute_comparison(all_sids)
            target_id = res['comparison_id']
            conn = get_db_connection()
            cursor = conn.cursor()

    # Load comparison record
    cursor.execute("SELECT * FROM comparisons WHERE id = ?", (target_id,))
    comp_row = cursor.fetchone()
    comp_dict = dict(comp_row)
    lcs_sequence = json.loads(comp_dict['lcs_sequence'])
    student_ids = json.loads(comp_dict['student_ids'])

    # Load comparison_students
    cursor.execute("""
        SELECT cs.*, s.branch
        FROM comparison_students cs
        JOIN students s ON cs.student_id = s.id
        WHERE cs.comparison_id = ?
        ORDER BY cs.id ASC
    """, (target_id,))
    comp_students = [dict(r) for r in cursor.fetchall()]

    # Load student activities
    for cs in comp_students:
        sid = cs['student_id']
        cursor.execute("SELECT * FROM activities WHERE student_id = ? ORDER BY order_index ASC, id ASC", (sid,))
        cs['activities'] = [dict(a) for a in cursor.fetchall()]
        miss_str = cs.get('missing_activities_list', '')
        cs['missing_list'] = [m.strip() for m in miss_str.split(',') if m.strip() and m.strip() != 'None']

    # Load missing activities table rows
    cursor.execute("SELECT * FROM missing_activities WHERE comparison_id = ? ORDER BY id ASC", (target_id,))
    missing_table_rows = [dict(r) for r in cursor.fetchall()]

    # Load category analysis
    cursor.execute("""
        SELECT category, COUNT(*) AS total_count
        FROM activities
        WHERE student_id IN ({})
        GROUP BY category
    """.format(','.join(['?'] * len(student_ids))), student_ids)
    cat_counts = {r['category']: r['total_count'] for r in cursor.fetchall()}

    conn.close()

    return {
        "comparison_id": target_id,
        "comparison_date": comp_dict['comparison_date'],
        "lcs_sequence": lcs_sequence,
        "lcs_length": comp_dict['lcs_length'],
        "student_ids": student_ids,
        "students": comp_students,
        "missing_table_rows": missing_table_rows,
        "category_counts": cat_counts
    }

def format_excel_sheet(ws, title_text=None):
    """Applies professional styling, borders, colors, auto filter, and column widths."""
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10, color="0F172A")
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Style Header row (row 1)
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[1].height = 28

    # Style Data rows
    for row in range(2, ws.max_row + 1):
        ws.row_dimensions[row].height = 20
        use_alt = (row % 2 == 0)
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = data_font
            cell.border = thin_border
            if use_alt:
                cell.fill = alt_fill
            # Align numbers / percentages center
            val_str = str(cell.value or '')
            if '%' in val_str or (isinstance(cell.value, (int, float)) and col > 3):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Freeze top row
    ws.freeze_panes = "A2"

    # Auto filter
    if ws.max_row > 1 and ws.max_column > 0:
        ws.auto_filter.ref = ws.dimensions

    # Auto column width
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or '')
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

def build_excel_workbook(comp_data):
    """
    Builds the 8 required Excel sheets following Sections 17, 18, 19, 20 to the letter.
    """
    wb = Workbook()
    # Remove default sheet
    default_sheet = wb.active

    comp_id = comp_data['comparison_id']
    comp_date = comp_data['comparison_date']
    lcs_sequence = comp_data['lcs_sequence']
    lcs_length = comp_data['lcs_length']
    students = comp_data['students']

    # --------------------------------------------------------------------------
    # 1. Sheet: Student Details
    # Columns: Name | Roll No | Branch | Total Activities | Common Activities | Missing Activities | Match %
    # --------------------------------------------------------------------------
    ws1 = wb.create_sheet(title="Student Details")
    ws1.append(["Name", "Roll No", "Branch", "Total Activities", "Common Activities", "Missing Activities", "Match %"])
    for s in students:
        miss_str = ", ".join(s['missing_list']) if s['missing_list'] else "None"
        ws1.append([
            s['student_name'],
            s['roll_no'],
            s['branch'],
            s['total_activities'],
            s['common_activities_count'],
            miss_str,
            f"{s['match_percentage']:.2f}%"
        ])
    format_excel_sheet(ws1)

    # --------------------------------------------------------------------------
    # 2. Sheet: All Activities (VERY IMPORTANT - SECTION 17 & 19)
    # EVERY ACTIVITY OF EVERY STUDENT MUST BE INCLUDED.
    # Columns: Name | Roll No | Branch | Activity Name | Common Activities | Total Activities | Missing Activities | Match %
    # --------------------------------------------------------------------------
    ws2 = wb.create_sheet(title="All Activities")
    ws2.append(["Name", "Roll No", "Branch", "Activity Name", "Common Activities", "Total Activities", "Missing Activities", "Match %"])
    for s in students:
        miss_str = ", ".join(s['missing_list']) if s['missing_list'] else "None"
        mp_str = f"{s['match_percentage']:.2f}%"
        for act in s['activities']:
            ws2.append([
                s['student_name'],
                s['roll_no'],
                s['branch'],
                act['activity_name'],
                s['common_activities_count'],
                s['total_activities'],
                miss_str,
                mp_str
            ])
    format_excel_sheet(ws2)

    # --------------------------------------------------------------------------
    # 3. Sheet: Comparison
    # Columns: Comparison ID | Date | Compared Students | LCS | LCS Length
    # --------------------------------------------------------------------------
    ws3 = wb.create_sheet(title="Comparison")
    ws3.append(["Comparison ID", "Date", "Compared Students", "LCS", "LCS Length"])
    compared_names = ", ".join([s['student_name'] for s in students])
    lcs_display = " -> ".join(lcs_sequence) if lcs_sequence else "No Common Sequence"
    ws3.append([comp_id, comp_date, compared_names, lcs_display, lcs_length])
    format_excel_sheet(ws3)

    # --------------------------------------------------------------------------
    # 4. Sheet: LCS Result
    # Columns: Comparison ID | Student | Common Activity | LCS Position
    # --------------------------------------------------------------------------
    ws4 = wb.create_sheet(title="LCS Result")
    ws4.append(["Comparison ID", "Student", "Common Activity", "LCS Position"])
    for pos, act_name in enumerate(lcs_sequence, 1):
        for s in students:
            ws4.append([comp_id, s['student_name'], act_name, pos])
    format_excel_sheet(ws4)

    # --------------------------------------------------------------------------
    # 5. Sheet: Missing Activities (SECTION 18 & 19)
    # Columns: Comparison ID | Student | Roll No | Missing Activity
    # --------------------------------------------------------------------------
    ws5 = wb.create_sheet(title="Missing Activities")
    ws5.append(["Comparison ID", "Student", "Roll No", "Missing Activity"])
    for s in students:
        if s['missing_list']:
            for miss in s['missing_list']:
                ws5.append([comp_id, s['student_name'], s['roll_no'], miss])
        else:
            ws5.append([comp_id, s['student_name'], s['roll_no'], "None (All Activities Common)"])
    format_excel_sheet(ws5)

    # --------------------------------------------------------------------------
    # 6. Sheet: Student Summary
    # Columns: Student | Roll No | Total Activities | Common Activities | Missing Activities | Match %
    # --------------------------------------------------------------------------
    ws6 = wb.create_sheet(title="Student Summary")
    ws6.append(["Student", "Roll No", "Total Activities", "Common Activities", "Missing Activities", "Match %"])
    for s in students:
        miss_str = ", ".join(s['missing_list']) if s['missing_list'] else "None"
        ws6.append([
            s['student_name'],
            s['roll_no'],
            s['total_activities'],
            s['common_activities_count'],
            miss_str,
            f"{s['match_percentage']:.2f}%"
        ])
    format_excel_sheet(ws6)

    # --------------------------------------------------------------------------
    # 7. Sheet: Activity-wise Percentage
    # Columns: Student | Activity | Common With Group | Activity Match % | Student Match %
    # --------------------------------------------------------------------------
    ws7 = wb.create_sheet(title="Activity-wise Percentage")
    ws7.append(["Student", "Activity", "Common With Group", "Activity Match %", "Student Match %"])
    for s in students:
        std_pct_str = f"{s['match_percentage']:.2f}%"
        # Determine which specific activity matched LCS in order
        act_names = [a['activity_name'].strip() for a in s['activities']]
        matched_set = set()
        l_idx = 0
        for i_act, a_name in enumerate(act_names):
            if l_idx < len(lcs_sequence) and a_name.lower() == lcs_sequence[l_idx].strip().lower():
                matched_set.add(i_act)
                l_idx += 1

        for i_act, act in enumerate(s['activities']):
            is_comm = (i_act in matched_set)
            comm_group_str = "Yes" if is_comm else "No"
            act_match_pct_str = "100%" if is_comm else "0%"
            ws7.append([
                s['student_name'],
                act['activity_name'],
                comm_group_str,
                act_match_pct_str,
                std_pct_str
            ])
    format_excel_sheet(ws7)

    # --------------------------------------------------------------------------
    # 8. Sheet: Category Analysis
    # Columns: Category | Total Activities | Common Activities | Percentage
    # --------------------------------------------------------------------------
    ws8 = wb.create_sheet(title="Category Analysis")
    ws8.append(["Category", "Total Activities", "Common Activities", "Percentage"])
    
    # Calculate category common activities
    # Fetch categories for common activities
    category_totals: dict[str, int] = {}
    category_commons: dict[str, int] = {}
    for s in students:
        act_names = [a['activity_name'].strip() for a in s['activities']]
        matched_set = set()
        l_idx = 0
        for i_act, a_name in enumerate(act_names):
            if l_idx < len(lcs_sequence) and a_name.lower() == lcs_sequence[l_idx].strip().lower():
                matched_set.add(i_act)
                l_idx += 1

        for i_act, act in enumerate(s['activities']):
            cat = act['category']
            category_totals[cat] = category_totals.get(cat, 0) + 1
            if i_act in matched_set:
                category_commons[cat] = category_commons.get(cat, 0) + 1

    for cat, tot in sorted(category_totals.items(), key=lambda x: x[0]):
        comm = category_commons.get(cat, 0)
        pct = (comm / tot * 100.0) if tot > 0 else 0.0
        ws8.append([cat, tot, comm, f"{pct:.2f}%"])

    format_excel_sheet(ws8)

    # Remove the initial default blank sheet
    if default_sheet in wb.worksheets:
        wb.remove(default_sheet)

    return wb

@app.route('/api/export/excel', methods=['GET'])
def export_excel():
    """Generates and serves professional Excel report (.xlsx)."""
    try:
        comparison_id_arg = request.args.get('comparison_id')
        comp_id = int(comparison_id_arg) if comparison_id_arg else None

        comp_data = get_or_run_active_comparison(comp_id)
        wb = build_excel_workbook(comp_data)

        timestamp_str = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"Longest_Common_Student_Activity_Comparison_{comp_data['comparison_id']}_{timestamp_str}.xlsx"
        filepath = os.path.join(REPORTS_EXCEL_DIR, filename)
        wb.save(filepath)

        return send_file(
            filepath,
            as_attachment=True,
            download_name=filename,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
    except Exception as e:
        return jsonify({"success": False, "error": f"Failed to generate Excel report: {str(e)}"}), 500

# ==============================================================================
# PDF REPORT GENERATION (reportlab) - SECTION 21
# ==============================================================================

class NumberedCanvas(canvas.Canvas):
    """Two-pass canvas for ReportLab to draw 'Page X of Y' in footer."""
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        getattr(self, '_startPage')()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count):
        self.saveState()
        # Top banner line
        self.setStrokeColor(colors.HexColor("#3B82F6"))
        self.setLineWidth(1.5)
        self.line(40, letter[1] - 40, letter[0] - 40, letter[1] - 40)

        # Footer line
        self.setStrokeColor(colors.HexColor("#CBD5E1"))
        self.setLineWidth(0.8)
        self.line(40, 45, letter[0] - 40, 45)

        # Footer text
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748B"))
        self.drawString(40, 32, "LONGEST COMMON STUDENT ACTIVITY | Dynamic Programming (LCS) System")
        page_num = getattr(self, '_pageNumber', 1)
        page_str = f"Page {page_num} of {page_count}"
        self.drawRightString(letter[0] - 40, 32, page_str)
        self.restoreState()

def build_pdf_document(comp_data, output_path):
    """Builds a beautiful, professional PDF report using ReportLab."""
    doc = SimpleDocTemplate(
        output_path,
        pagesize=letter,
        leftMargin=40,
        rightMargin=40,
        topMargin=55,
        bottomMargin=55
    )

    styles = getSampleStyleSheet()
    primary_color = colors.HexColor("#1E293B")
    accent_blue = colors.HexColor("#2563EB")
    slate_dark = colors.HexColor("#0F172A")

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=primary_color,
        alignment=1, # Center
        spaceAfter=4
    )

    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=15,
        textColor=accent_blue,
        alignment=1, # Center
        spaceAfter=14
    )

    h2_style = ParagraphStyle(
        'H2',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=17,
        textColor=primary_color,
        spaceBefore=12,
        spaceAfter=6
    )

    body_style = ParagraphStyle(
        'Body',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=slate_dark
    )

    lcs_box_style = ParagraphStyle(
        'LcsBox',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#1E3A8A"),
        alignment=1
    )

    table_header_style = ParagraphStyle(
        'TH',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=11,
        textColor=colors.white,
        alignment=1
    )

    table_cell_style = ParagraphStyle(
        'TD',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        textColor=slate_dark
    )

    table_cell_center = ParagraphStyle(
        'TDC',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        textColor=slate_dark,
        alignment=1
    )

    story: list[Any] = []

    # Title & Subtitle
    story.append(Paragraph("Longest Common Student Activity Report", title_style))
    story.append(Paragraph("DAA Hackathon Project | Dynamic Programming LCS Comparison & Analytics", subtitle_style))
    story.append(Spacer(1, 4))

    # Metadata Summary Bar
    comp_id = comp_data['comparison_id']
    comp_date = comp_data['comparison_date']
    lcs_sequence = comp_data['lcs_sequence']
    lcs_length = comp_data['lcs_length']
    students = comp_data['students']

    meta_table_data = [
        [
            Paragraph(f"<b>Comparison ID:</b> #{comp_id}", body_style),
            Paragraph(f"<b>Generated At:</b> {comp_date}", body_style),
            Paragraph(f"<b>Total Students:</b> {len(students)}", body_style),
            Paragraph(f"<b>LCS Length:</b> {lcs_length}", body_style)
        ]
    ]
    meta_table = Table(meta_table_data, colWidths=[110, 160, 120, 142])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#F1F5F9")),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor("#CBD5E1")),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 6),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ('LEFTPADDING', (0,0), (-1,-1), 8),
        ('RIGHTPADDING', (0,0), (-1,-1), 8),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 10))

    # Common Activity Sequence Callout Box
    story.append(Paragraph("Common Activity Sequence (LCS via Dynamic Programming)", h2_style))
    lcs_text = " &nbsp;➔&nbsp; ".join([f"<font color='#2563EB'><b>{act}</b></font>" for act in lcs_sequence]) if lcs_sequence else "No common subsequence found"
    lcs_box_data = [[
        Paragraph(f"<b>Discovered Common Sequence (Length = {lcs_length}):</b><br/>{lcs_text}", lcs_box_style)
    ]]
    lcs_box_table = Table(lcs_box_data, colWidths=[532])
    lcs_box_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#EFF6FF")),
        ('BOX', (0,0), (-1,-1), 1.5, colors.HexColor("#3B82F6")),
        ('TOPPADDING', (0,0), (-1,-1), 10),
        ('BOTTOMPADDING', (0,0), (-1,-1), 10),
        ('LEFTPADDING', (0,0), (-1,-1), 12),
        ('RIGHTPADDING', (0,0), (-1,-1), 12),
    ]))
    story.append(lcs_box_table)
    story.append(Spacer(1, 12))

    # Section: Compared Students & Performance Summary Table
    story.append(Paragraph("Student Performance & Match % Summary", h2_style))
    story.append(Paragraph(
        "<i>Match % Formula: (Common Activities / Total Activities of that student) × 100</i>", 
        ParagraphStyle('Note', parent=styles['Normal'], fontSize=8, leading=10, textColor=colors.HexColor("#64748B"), spaceAfter=6)
    ))

    summary_headers = [
        Paragraph("Student Name", table_header_style),
        Paragraph("Roll No", table_header_style),
        Paragraph("Branch", table_header_style),
        Paragraph("Total Acts", table_header_style),
        Paragraph("Common", table_header_style),
        Paragraph("Missing", table_header_style),
        Paragraph("Match %", table_header_style)
    ]
    summary_rows = [summary_headers]

    for s in students:
        miss_count = len(s['missing_list'])
        match_color = "#059669" if s['match_percentage'] >= 70 else ("#D97706" if s['match_percentage'] >= 40 else "#DC2626")
        summary_rows.append([
            Paragraph(f"<b>{s['student_name']}</b>", table_cell_style),
            Paragraph(s['roll_no'], table_cell_center),
            Paragraph(s['branch'], table_cell_style),
            Paragraph(str(s['total_activities']), table_cell_center),
            Paragraph(f"<b>{s['common_activities_count']}</b>", table_cell_center),
            Paragraph(str(miss_count), table_cell_center),
            Paragraph(f"<b><font color='{match_color}'>{s['match_percentage']:.2f}%</font></b>", table_cell_center)
        ])

    sum_table = Table(summary_rows, colWidths=[110, 55, 127, 60, 60, 60, 60])
    sum_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#1E293B")),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E1")),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor("#F8FAFC")]),
        ('TOPPADDING', (0,0), (-1,-1), 5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
    ]))
    story.append(sum_table)
    story.append(Spacer(1, 14))

    # Section: Missing Activities Breakdown
    story.append(Paragraph("Missing Activities Breakdown (Per Student)", h2_style))
    miss_box_rows = [
        [
            Paragraph("Student Name", table_header_style),
            Paragraph("Roll No", table_header_style),
            Paragraph("Missing Activities (Not in LCS)", table_header_style)
        ]
    ]
    for s in students:
        miss_display = ", ".join(s['missing_list']) if s['missing_list'] else "<font color='#059669'>None (All Activities Common)</font>"
        miss_box_rows.append([
            Paragraph(f"<b>{s['student_name']}</b>", table_cell_style),
            Paragraph(s['roll_no'], table_cell_center),
            Paragraph(miss_display, table_cell_style)
        ])

    miss_table = Table(miss_box_rows, colWidths=[120, 60, 352])
    miss_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#334155")),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E1")),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor("#F8FAFC")]),
        ('TOPPADDING', (0,0), (-1,-1), 5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
    ]))
    story.append(miss_table)
    story.append(Spacer(1, 14))

    # Section: Complete Activities Master List
    story.append(Paragraph("Master Activity Log (All Students & Activities Included)", h2_style))
    act_headers = [
        Paragraph("Student", table_header_style),
        Paragraph("Activity Name", table_header_style),
        Paragraph("Category", table_header_style),
        Paragraph("Date", table_header_style),
        Paragraph("Status", table_header_style),
        Paragraph("In LCS?", table_header_style)
    ]
    act_rows = [act_headers]

    for s in students:
        act_names = [a['activity_name'].strip() for a in s['activities']]
        matched_set = set()
        l_idx = 0
        for i_act, a_name in enumerate(act_names):
            if l_idx < len(lcs_sequence) and a_name.lower() == lcs_sequence[l_idx].strip().lower():
                matched_set.add(i_act)
                l_idx += 1

        for i_act, act in enumerate(s['activities']):
            is_comm = (i_act in matched_set)
            lcs_badge = "<font color='#059669'><b>Yes (LCS)</b></font>" if is_comm else "<font color='#DC2626'>Missing</font>"
            act_rows.append([
                Paragraph(s['student_name'], table_cell_style),
                Paragraph(f"<b>{act['activity_name']}</b>", table_cell_style),
                Paragraph(act['category'], table_cell_center),
                Paragraph(act['activity_date'], table_cell_center),
                Paragraph(act['verification_status'], table_cell_center),
                Paragraph(lcs_badge, table_cell_center)
            ])

    act_table = Table(act_rows, colWidths=[90, 150, 75, 75, 75, 67])
    act_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#1E293B")),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E1")),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor("#F8FAFC")]),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
    ]))
    story.append(act_table)

    doc.build(story, canvasmaker=NumberedCanvas)

@app.route('/api/export/pdf', methods=['GET'])
def export_pdf():
    """Generates and serves professional PDF report (.pdf)."""
    try:
        comparison_id_arg = request.args.get('comparison_id')
        comp_id = int(comparison_id_arg) if comparison_id_arg else None

        comp_data = get_or_run_active_comparison(comp_id)

        timestamp_str = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"Longest_Common_Student_Activity_Report_{comp_data['comparison_id']}_{timestamp_str}.pdf"
        filepath = os.path.join(REPORTS_PDF_DIR, filename)

        build_pdf_document(comp_data, filepath)

        return send_file(
            filepath,
            as_attachment=True,
            download_name=filename,
            mimetype="application/pdf"
        )
    except Exception as e:
        return jsonify({"success": False, "error": f"Failed to generate PDF report: {str(e)}"}), 500

# ==============================================================================
# MAIN ENTRYPOINT
# ==============================================================================

if __name__ == '__main__':
    # Running locally on http://127.0.0.1:5000/
    print("==================================================================")
    print(" LONGEST COMMON STUDENT ACTIVITY (DAA Hackathon Project) ")
    print(" Running at: http://127.0.0.1:5000/")
    print(" Single Page Application with Dynamic Programming LCS Algorithm")
    print("==================================================================")
    app.run(host='127.0.0.1', port=5000, debug=True)
